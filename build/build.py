#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The only build command.

    uv run python build/build.py            what the model currently says
    uv run python build/build.py check      validate facts/ and scales/, report problems
    uv run python build/build.py workbook   write out/AI-Capability-Model.xlsx
    uv run python build/build.py views      write the reports and views to out/
    uv run python build/build.py all        check, then workbook + views
    uv run python build/build.py ingest [path ...]   read reviewer edits back into facts/
    uv run python build/build.py test       run the test suite

facts/ is edited. out/ is generated. `ingest` is the one path that writes to
facts/, and it only ever writes observations.
"""
import json, os, re, sys, collections

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import facts as F

OUT = os.path.join(ROOT, 'out')
REVIEW = os.path.join(ROOT, 'review')
WORKBOOK = os.path.join(OUT, 'AI-Capability-Model.xlsx')
OBS_SHEET = "2. Observations"


def _model():
    return F.Model()


# --------------------------------------------------------------- status
def cmd_status():
    m = _model()
    scales = F.load_scales()
    print("facts")
    print("  %-24s %d domains, %d capabilities, %d criteria"
          % ("capabilities.json", len(m.domains), len(m.capabilities),
             sum(len(c['criteria']) for c in m.capabilities)))
    print("  %-24s %d offerings, %d in-the-box questions"
          % ("offerings.json", len(m.offerings),
             sum(len(o['in_the_box']) for o in m.offerings)))
    print("  %-24s %d assets, %d released, %d statuses declared"
          % ("assets.json", len(m.assets),
             sum(1 for a in m.assets if m.released(a)), len(m.asset_statuses)))
    print("  %-24s %d units, %d mapped" % ("owners.json", len(m.units),
                                           len(m.owner_map)))
    print("  %-24s %d rows" % ("observations.json", len(m.observations)))
    gr = collections.Counter(s['grade'] for s in m.sources)
    print("  %-24s %d sources (%s)" % ("sources.json", len(m.sources),
          ", ".join("%s %d" % (g, gr[g]) for g in sorted(gr))))
    print("  %-24s %d candidate references" % ("obligations.json", len(m.obligations)))
    print("  %-24s %d use-case question(s): %s" % ("questions.json", len(m.questions),
          ", ".join(q['id'] for q in m.questions)))

    print("\nobservations, by capability")
    vals = {c['id']: m.values(c['id']) for c in m.capabilities}
    for t in m.observation_types:
        c = collections.Counter(vals[x['id']][t['id']] for x in m.capabilities)
        derived = "  (derived from L3)" if t['id'] in m.criterion_types else ""
        print("  %-10s %s%s" % (t['id'], dict(sorted(c.items())), derived))

    for t in m.criterion_types:
        c = collections.Counter(
            r.get('value', 'unknown')
            for x in m.capabilities for _, r in m.criteria_obs(x['id'], t))
        n = sum(len(x['criteria']) for x in m.capabilities)
        print("\n%s, by L3 criterion (%d rows - this is what reviewers answer)"
              % (t, n))
        print("  %-10s %s" % (t, dict(sorted(c.items()))))

    print("\nscales")
    for s in scales:
        r = m.rate_all(s)
        c = collections.Counter(v[0] for v in r.values())
        pretty = {("not rated" if k is None else k): n
                  for k, n in sorted(c.items(), key=lambda kv: (kv[0] is None, kv[0]))}
        star = "  (default)" if getattr(s, 'DEFAULT', False) else "  (lens)"
        print("  %-24s %s%s" % (s.NAME, pretty, star))

    nomatch = [c['id'] for c in m.capabilities if m.owner(c['id'])[1] == 'NO MATCH']
    if nomatch:
        print("\n%d capabilities have no owner in the Bank's catalogue: %s"
              % (len(nomatch), ", ".join(nomatch)))
    return 0


# --------------------------------------------------------------- check
def check(m=None, scales=None):
    """(problems, advisories) for facts/ and scales/.

    A problem is something a view would render wrongly or a rule of the model
    would be broken by - `check` fails on any.  An advisory is worth a look but
    does not stop a build: a register that moved on since an observation was
    taken, a value recorded without a date.
    """
    m = m or _model()
    scales = scales if scales is not None else F.load_scales()
    problems, advisories = [], []
    ids = {c['id'] for c in m.capabilities}
    crit = {x['id'] for c in m.capabilities for x in c['criteria']}
    crit_of = m.criterion_parent

    # ---- capability map
    seen = collections.Counter(c['id'] for c in m.capabilities)
    for cid, n in seen.items():
        if n > 1:
            problems.append("capability id %s appears %d times" % (cid, n))
    for c in m.capabilities:
        if c['domain'] not in m.domain_by_id:
            problems.append("%s is in unknown domain %s" % (c['id'], c['domain']))
        for x in c['criteria']:
            if not x['id'].startswith(c['id'] + '.'):
                problems.append("criterion %s is filed under %s" % (x['id'], c['id']))

    # ---- offerings and assets
    for o in m.offerings:
        for cid in o['enables']:
            if cid not in ids:
                problems.append("%s enables unknown capability %s" % (o['id'], cid))
        for a in o['assets']:
            if a not in m.asset_by_id:
                problems.append("%s references unknown asset %s" % (o['id'], a))
        for x in o['in_the_box']:
            # controls point at an L3 criterion, which resolves to its L2
            if x['capability'] not in ids and x['capability'] not in crit:
                problems.append("%s in-the-box control points at unknown subject %s"
                                % (o['id'], x['capability']))
        for k in ('assets_released', 'assets_total'):
            if k in o:
                problems.append("%s stores %s, which is derived from asset statuses - "
                                "remove it" % (o['id'], k))
    for a in m.assets:
        if a['status'] not in m.asset_statuses:
            problems.append("asset %s has status %r, not declared in assets.json "
                            "`statuses`" % (a['id'], a['status']))
    for st, d in m.asset_statuses.items():
        if not isinstance(d, dict) or 'released' not in d:
            problems.append("status %r must declare `released: true|false`" % st)

    # ---- observations
    if set(m.observation_values) != set(F.OBS_VALUES):
        problems.append("observations.json declares values %s but F.OBS_VALUES "
                        "says %s - the vocabulary must have one definition"
                        % (list(m.observation_values), list(F.OBS_VALUES)))
    for v in m.observation_values:
        if v not in F.OBS_MEANING:
            problems.append("value %r is accepted but says nowhere what it means, so "
                            "no view can explain it - add it to F.OBS_MEANING" % v)
    for r in m.observations:
        where = "%s/%s" % (r.get('criterion') or r['capability'], r['observation'])
        if r['capability'] not in ids:
            problems.append("observation for unknown capability %s" % r['capability'])
        if r.get('criterion'):
            if r['criterion'] not in crit:
                problems.append("observation for unknown criterion %s" % r['criterion'])
            elif crit_of[r['criterion']] != r['capability']:
                problems.append("%s is recorded under %s but belongs to %s"
                                % (r['criterion'], r['capability'],
                                   crit_of[r['criterion']]))
        if r['observation'] not in {t['id'] for t in m.observation_types}:
            problems.append("%s has unknown observation type" % where)
        if r['value'] not in m.observation_values:
            problems.append("%s has invalid value %r" % (where, r['value']))
        if r['value'] == 'n/a' and not r.get('basis'):
            problems.append("%s is n/a with no reason" % where)
        if r['value'] in ('yes', 'partial') and not r.get('evidence'):
            problems.append("%s claims %s with no evidence" % (where, r['value']))
        if r['value'] == 'no' and not (r.get('evidence') or r.get('basis')):
            problems.append("%s claims no with neither evidence nor basis - a `no` is "
                            "an evidenced negative, not an absence" % where)
        d = r.get('observed_on') or ''
        if d and not re.fullmatch(r'\d{4}-\d{2}-\d{2}', d):
            problems.append("%s has observed_on %r, expected YYYY-MM-DD" % (where, d))
        if r['value'] != 'unknown' and not (d and r.get('observed_by')):
            advisories.append("%s is %s but has no observer or date" % (where, r['value']))

    # Capability-level observations: one row per capability per type.
    # Criterion-level observations (ADR-0014): one row per criterion.
    cap_types = {t['id'] for t in m.observation_types
                 if t['id'] not in m.criterion_types}
    for c in m.capabilities:
        got = set(dict.get(m.obs_by_cap, c['id'], {}))
        for t in cap_types - got:
            problems.append("%s has no %s observation" % (c['id'], t))
        for t in m.criterion_types:
            rows = m.obs_by_crit.get(c['id'], {}).get(t, {})
            missing = [x['id'] for x in c['criteria'] if x['id'] not in rows]
            if missing:
                problems.append("%s has no %s observation for %d criteria (%s%s)"
                                % (c['id'], t, len(missing), ", ".join(missing[:3]),
                                   ", ..." if len(missing) > 3 else ""))
            if any(k in m.obs_by_crit.get(c['id'], {}).get(t, {})
                   for k in ()):  # placeholder for future per-criterion rules
                pass
        for t in m.criterion_types:
            if t in got:
                problems.append("%s has a capability-level %s row; %s is recorded per "
                                "criterion (ADR-0014)" % (c['id'], t, t))

    # ---- the reason register behind `enabled: n/a`
    ctx = m.enablement_context
    ctx_caps = set(ctx.get('realized_by_enterprise_service', {})) | \
        set(ctx.get('purely_organizational', []))
    for cid in ctx_caps:
        if cid not in ids:
            problems.append("enablement-context names unknown capability %s" % cid)
    for c in m.capabilities:
        v = m.cap_obs(c['id'], 'enabled').get('value')
        if v == 'n/a' and c['id'] not in ctx_caps:
            problems.append("%s/enabled is n/a but enablement-context.json records no "
                            "reason for it" % c['id'])
        if c['id'] in ctx_caps and c['id'] in m.offerings_for:
            advisories.append("%s is in enablement-context (no platform offering expected) "
                              "but %s enables it"
                              % (c['id'], ", ".join(o['id'] for o in m.offerings_for[c['id']])))
        if c['id'] in ctx_caps and v not in ('n/a', None) and c['id'] not in m.offerings_for:
            advisories.append("%s is in enablement-context but its enabled observation is %r, "
                              "not n/a" % (c['id'], v))

    # ---- register drift: observations seeded from the register that the
    # register no longer supports (an asset was released, an offering changed)
    for c in m.capabilities:
        implied = m.register_implies(c['id'])
        for t in ('enabled', 'defined'):
            row = m.cap_obs(c['id'], t)
            if 'asset register' in (row.get('observed_by') or '') \
                    and row.get('value') != implied[t]:
                advisories.append("%s/%s was recorded as %s from the asset register; the "
                                  "register now implies %s - re-observe or confirm"
                                  % (c['id'], t, row.get('value'), implied[t]))

    # ---- owners
    for cid in ids:
        if cid not in m.owner_map:
            problems.append("%s is not mapped to a Bank unit" % cid)
    for cid, mp in m.owner_map.items():
        if cid not in ids:
            problems.append("owners.json maps unknown capability %s" % cid)
        if mp.get('unit') and mp['unit'] not in m.unit_by_name:
            problems.append("%s is mapped to unit %r, not in the units list"
                            % (cid, mp['unit']))
        if not mp.get('unit') and mp.get('match') != 'NO MATCH':
            problems.append("%s has no unit but match is %r, expected NO MATCH"
                            % (cid, mp.get('match')))

    # ---- sources (ADR-0010)
    for c in m.capabilities:
        for cit, src, _loc in m.capability_sources(c['id']):
            if src is None:
                problems.append("%s cites %r, which does not resolve in sources.json"
                                % (c['id'], cit))
    for s in m.sources:
        if s.get('grade') not in m.source_grades:
            problems.append("source %s has grade %r, not one of %s"
                            % (s['id'], s.get('grade'), sorted(m.source_grades)))
    for cit, e in m._normalize.items():
        if e.get('source') not in m.source_by_id:
            problems.append("normalize entry %r points at unknown source %r"
                            % (cit, e.get('source')))

    # ---- obligations (ADR-0008)
    for ob in m.obligations:
        if ob['capability'] not in ids:
            problems.append("obligation %r names unknown capability %s"
                            % (ob['instrument'], ob['capability']))

    # ---- questions
    outs = collections.Counter()
    for q in m.questions:
        for k in ('id', 'output', 'title'):
            if not q.get(k):
                problems.append("question %r lacks %s" % (q.get('id', '?'), k))
        if not re.fullmatch(r'[a-z0-9][a-z0-9_-]*', q.get('output', '')):
            problems.append("question %s output %r must be a safe file name"
                            % (q.get('id'), q.get('output')))
        outs[q.get('output')] += 1
        for x in q.get('offerings', []):
            if x not in m.offering_by_id:
                problems.append("question %s names unknown offering %s" % (q['id'], x))
        for x in q.get('capabilities', []):
            if x not in ids:
                problems.append("question %s names unknown capability %s" % (q['id'], x))
        if q.get('finding') and not (q.get('finding_by') and q.get('finding_on')):
            problems.append("question %s has a finding with no finding_by / finding_on - "
                            "a judgement carries who made it and when" % q['id'])
        for step in q.get('next', []):
            if not (step.get('who') and step.get('what')):
                problems.append("question %s has a next step without who/what" % q['id'])
    for name, n in outs.items():
        if n > 1:
            problems.append("%d questions share the output name %r" % (n, name))
    reserved = {'capability-assessment-%s' % s.SHORT for s in scales} | \
        {'management-report', 'management-report-illustrative', 'provenance'}
    for q in m.questions:
        if q.get('output') in reserved:
            problems.append("question %s output %r collides with a generated view"
                            % (q['id'], q['output']))

    # ---- scales
    problems.extend(F.validate_scales(scales))
    return problems, advisories


def cmd_check():
    m = _model()
    problems, advisories = check(m)
    if problems:
        print("%d problem(s):" % len(problems))
        for p in problems:
            print("  -", p)
    else:
        print("facts/ is consistent: %d capabilities, %d observations, %d offerings, "
              "%d assets, %d sources, %d question(s); %d scale(s) sound"
              % (len(m.capabilities), len(m.observations), len(m.offerings),
                 len(m.assets), len(m.sources), len(m.questions), len(F.load_scales())))
    if advisories:
        print("%d advisory(ies) - worth a look, not a failure:" % len(advisories))
        for a in advisories[:25]:
            print("  ~", a)
        if len(advisories) > 25:
            print("  ~ ... and %d more" % (len(advisories) - 25))
    return 1 if problems else 0


# --------------------------------------------------------------- workbook
def cmd_workbook():
    import build_workbook
    m = _model()
    os.makedirs(OUT, exist_ok=True)
    sheets = build_workbook.build(m, F.default_scale(), WORKBOOK)
    print("wrote %s" % os.path.relpath(WORKBOOK, ROOT))
    print("  sheets: %s" % ", ".join(sheets))
    return 0


# --------------------------------------------------------------- views
def build_views_to(m, out):
    """Write every view to `out`. Returns the file names written, in order.

    The capability map on its own; one capability page per scale; the
    management report as text and as a self-contained page, plus the
    illustrative edition; one page per use-case question; the provenance view;
    the walkthrough deck.  Nothing here names a scale or a question.
    """
    import build_views, build_report, build_deck
    os.makedirs(out, exist_ok=True)
    written = []

    def emit(name, text, label=None):
        path = os.path.join(out, name)
        with open(path, 'w', encoding='utf-8') as f:
            f.write(text)
        written.append(label or name)

    emit('capability-map.md', build_views.capability_map(m),
         'capability-map.md  (the taxonomy alone)')
    scales = F.load_scales()
    default = F.default_scale()
    for s in scales:
        emit('capability-assessment-%s.md' % s.SHORT, build_views.capability_view(m, s),
             'capability-assessment-%s.md%s' % (s.SHORT, '  (default scale)'
                                                if s is default else '  (lens)'))
    emit('management-report.md', build_views.management_report(m, default))
    emit('management-report.html', build_report.report(m, default, scales=scales))
    emit('management-report-illustrative.html',
         build_report.report(m, default, demo=True, scales=scales),
         'management-report-illustrative.html  (SAMPLE observations)')
    for q in m.questions:
        emit('%s.md' % q['output'], build_views.question_view(m, q),
             '%s.md  (question: %s)' % (q['output'], q['id']))
    emit('provenance.md', build_views.provenance_view(m, scales))
    emit('walkthrough.html', build_deck.deck(m, default, scales=scales),
         'walkthrough.html  (the deck)')
    return written


def cmd_views():
    m = _model()
    written = build_views_to(m, OUT)
    print("wrote %d view(s) to out/:" % len(written))
    for x in written:
        print("  -", x)
    return 0


# --------------------------------------------------------------- ingest
def ingest_workbook(path, doc):
    """Apply sheet 2 of one returned workbook to an observations document.

    Pure with respect to the file system: `doc` is modified in place and the
    caller decides whether to write it.  Returns (changed, skipped) where
    `changed` is a list of (key, old_value, new_value) and `skipped` a list of
    human-readable reasons rows were ignored.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    if OBS_SHEET not in wb.sheetnames:
        raise ValueError("%s has no '%s' sheet" % (path, OBS_SHEET))
    ws = wb[OBS_SHEET]
    # key: (capability, observation, criterion or '') - criterion-level rows
    # (ADR-0014) are addressed by their criterion id
    index = {(r['capability'], r['observation'], r.get('criterion') or ''): r
             for r in doc['observations']}
    changed, skipped = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        # sheet 2 columns: capability, name, observation, criterion, subject,
        # question, VALUE, evidence, observed_by, date, basis
        (cid, _nm, typ, crit, _subj, _q,
         value, evidence, by, when) = (list(row) + [None] * 10)[:10]
        if typ is None:
            continue          # capability banner row, or the footer note
        key = (str(cid).strip(), str(typ).strip(),
               str(crit).strip() if crit else '')
        rec = index.get(key)
        if rec is None:
            skipped.append("unknown row " + "/".join(x for x in key if x))
            continue
        v = (str(value).strip().lower().replace('n-a', 'n/a')
             if value is not None else 'unknown')
        if v not in doc['values']:
            skipped.append("%s/%s has value %r" % (key[0], key[1], value))
            continue
        new = {'value': v,
               'evidence': (evidence or '').strip() if evidence else '',
               'observed_by': (by or '').strip() if by else '',
               'observed_on': str(when).strip()[:10] if when else ''}
        if any(rec.get(k, '') != v_ for k, v_ in new.items()):
            changed.append((key, rec.get('value'), v))
            rec.update(new)
    return changed, skipped


def cmd_ingest(paths=None):
    """Read sheet 2 of one or more returned workbooks back into
    facts/observations.json.

    Several reviewers return several files.  Each is applied in the order
    given; a cell changed by more than one file is reported, because the last
    file wins and someone should know that it did.
    """
    paths = list(paths or [])
    if not paths:
        paths = [WORKBOOK]
    for p in paths:
        if not os.path.exists(p):
            print("no workbook at %s - run 'workbook' first, or pass a path" % p)
            return 1
    doc = json.load(open(os.path.join(F.FACTS, 'observations.json'), encoding='utf-8'))
    touched = {}
    total = 0
    for p in paths:
        try:
            changed, skipped = ingest_workbook(p, doc)
        except ValueError as e:
            print(e)
            return 1
        print("ingested %s: %d observation(s) updated"
              % (os.path.relpath(p, ROOT) if p.startswith(ROOT) else p, len(changed)))
        for key, old, new in changed:
            prev = touched.get(key)
            if prev:
                print("  ! %s was also set by %s (%s); this file sets %s"
                      % ("/".join(x for x in key if x), prev[0], prev[1], new))
            touched[key] = (os.path.basename(p), new)
        for u in skipped[:10]:
            print("  skipped:", u)
        if len(skipped) > 10:
            print("  ... and %d more" % (len(skipped) - 10))
        total += len(changed)
    F.save_facts('observations.json', doc)
    if total:
        print("now run: uv run python build/build.py all")
    return 0


# --------------------------------------------------------------- test
def cmd_test():
    import unittest
    suite = unittest.defaultTestLoader.discover(os.path.join(ROOT, 'tests'),
                                                top_level_dir=ROOT)
    result = unittest.TextTestRunner(verbosity=1).run(suite)
    return 0 if result.wasSuccessful() else 1


CMDS = {'status': cmd_status, 'check': cmd_check, 'workbook': cmd_workbook,
        'views': cmd_views, 'test': cmd_test}


def main(argv):
    if not argv:
        return cmd_status()
    cmd = argv[0]
    if cmd == 'all':
        # never build views from facts that do not pass: the numbers would be wrong
        if cmd_check():
            print("\nnot building: fix the problems above first")
            return 1
        return cmd_workbook() or cmd_views()
    if cmd == 'ingest':
        return cmd_ingest(argv[1:])
    if cmd not in CMDS:
        print(__doc__)
        print("unknown command: %s" % cmd)
        return 1
    return CMDS[cmd]() or 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
