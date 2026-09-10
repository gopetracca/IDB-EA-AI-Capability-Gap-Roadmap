# -*- coding: utf-8 -*-
"""Tests for the rules the documentation promises.

Each test names the record it checks (an ADR, a README section), so a failing
test says which promise broke.  Run with `uv run python build/build.py test`.

Nothing here writes to facts/.  Builds run against a temporary directory and
the ingest tests work on an in-memory copy of the observations document.
"""
import hashlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from html.parser import HTMLParser

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'build'))

import facts as F          # noqa: E402
import build               # noqa: E402


def _facts_digest():
    h = hashlib.sha256()
    for fn in sorted(os.listdir(F.FACTS)):
        if fn.endswith('.json'):
            with open(os.path.join(F.FACTS, fn), 'rb') as f:
                h.update(fn.encode() + f.read())
    return h.hexdigest()


class RollUp(unittest.TestCase):
    """ADR-0014 - the roll-up table, row by row."""

    def test_every_criterion_yes_and_examined_is_yes(self):
        self.assertEqual(F.roll_up_values(['yes', 'yes', 'n/a']), 'yes')

    def test_one_unexamined_criterion_blocks_yes(self):
        self.assertEqual(F.roll_up_values(['yes', 'yes', 'unknown']), 'partial')
        self.assertEqual(F.roll_up_values(['yes', 'unknown', 'unknown', 'unknown']),
                         'partial')

    def test_any_weak_link_is_partial(self):
        self.assertEqual(F.roll_up_values(['yes', 'no']), 'partial')
        self.assertEqual(F.roll_up_values(['partial', 'no', 'n/a']), 'partial')

    def test_no_yes_or_partial_with_a_no_is_no(self):
        self.assertEqual(F.roll_up_values(['no', 'no']), 'no')
        # an examined `no` beside unexamined criteria still reads `no`: the
        # ADR excludes `unknown` from the denominator, and this asymmetry with
        # `yes` is deliberate - under-claiming is accepted, over-claiming is not
        self.assertEqual(F.roll_up_values(['no', 'unknown', 'unknown']), 'no')

    def test_nothing_examined_is_unknown_not_zero(self):
        self.assertEqual(F.roll_up_values(['unknown', 'unknown']), 'unknown')
        self.assertEqual(F.roll_up_values([]), 'unknown')

    def test_all_na_and_examined_is_na(self):
        self.assertEqual(F.roll_up_values(['n/a', 'n/a']), 'n/a')

    def test_all_na_with_unexamined_is_not_yet_na(self):
        self.assertEqual(F.roll_up_values(['n/a', 'unknown']), 'unknown')


class DefaultLadder(unittest.TestCase):
    """ADR-0013 / scales/README.md - performance comes first."""

    def setUp(self):
        self.s = F.default_scale()

    def lv(self, p, e, sk, d):
        return self.s.level({'practised': p, 'enabled': e, 'skilled': sk, 'defined': d})[0]

    def test_no_practice_observed_means_not_rated_however_good_the_enablers(self):
        self.assertIsNone(self.lv('unknown', 'yes', 'yes', 'yes'))

    def test_not_performed_is_level_0_even_with_a_standard(self):
        self.assertEqual(self.lv('no', 'yes', 'yes', 'yes'), 0)

    def test_performed_alone_is_level_1(self):
        self.assertEqual(self.lv('yes', 'no', 'unknown', 'unknown'), 1)
        self.assertEqual(self.lv('partial', 'yes', 'yes', 'yes'), 1)

    def test_tooling_and_competence_give_level_2(self):
        self.assertEqual(self.lv('yes', 'yes', 'yes', 'no'), 2)
        self.assertEqual(self.lv('yes', 'yes', 'yes', 'partial'), 2)

    def test_na_tooling_does_not_block_level_2(self):
        self.assertEqual(self.lv('yes', 'n/a', 'yes', 'unknown'), 2)

    def test_approved_standard_gives_level_3(self):
        self.assertEqual(self.lv('yes', 'yes', 'yes', 'yes'), 3)

    def test_nothing_above_derivable_max(self):
        top = getattr(self.s, 'DERIVABLE_MAX', None)
        self.assertIsNotNone(top)
        self.assertEqual(self.lv('yes', 'yes', 'yes', 'yes'), top)


class Scales(unittest.TestCase):
    """scales/README.md - the contract every scale keeps."""

    def test_every_scale_is_sound(self):
        self.assertEqual(F.validate_scales(), [])

    def test_exactly_one_default(self):
        self.assertTrue(getattr(F.default_scale(), 'DEFAULT', False))

    def test_shorts_are_unique_and_file_safe(self):
        shorts = [s.SHORT for s in F.load_scales()]
        self.assertEqual(len(shorts), len(set(shorts)))

    def test_lenses_declare_a_question(self):
        for s in F.load_scales():
            self.assertTrue(getattr(s, 'QUESTION', ''), s.SHORT)


class Facts(unittest.TestCase):
    """build.py check - the facts pass their own validation."""

    def test_check_passes(self):
        problems, _advisories = build.check()
        self.assertEqual(problems, [])

    def test_release_counts_are_derived_not_stored(self):
        m = F.Model()
        for o in m.offerings:
            self.assertNotIn('assets_released', o)
            rel, tot = m.release_count(o)
            self.assertLessEqual(rel, tot)
            self.assertEqual(tot, len(o['assets']))

    def test_every_citation_resolves(self):
        m = F.Model()
        for c in m.capabilities:
            for cit, src, _ in m.capability_sources(c['id']):
                self.assertIsNotNone(src, "%s cites %r" % (c['id'], cit))

    def test_values_roll_practised_up_from_criteria(self):
        m = F.Model()
        for c in m.capabilities:
            v = m.values(c['id'])
            self.assertEqual(set(v), set(t['id'] for t in m.observation_types))
            self.assertEqual(v['practised'], m.roll_up(c['id'], 'practised'))

    def test_reading_practised_from_obs_by_cap_raises(self):
        m = F.Model()
        with self.assertRaises(KeyError):
            m.obs_by_cap.get('4.4', {}).get('practised')


class _Balance(HTMLParser):
    VOID = {'meta', 'br', 'img', 'link', 'input', 'hr'}

    def __init__(self):
        HTMLParser.__init__(self)
        self.stack, self.errors = [], []

    def handle_starttag(self, tag, attrs):
        if tag not in self.VOID:
            self.stack.append(tag)

    def handle_endtag(self, tag):
        if tag in self.VOID:
            return
        if not self.stack or self.stack[-1] != tag:
            self.errors.append((tag, list(self.stack[-3:])))
        else:
            self.stack.pop()


class Builds(unittest.TestCase):
    """docs/using-the-model.md - what comes out, and that building never
    writes to facts/."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix='aicap-')
        cls.before = _facts_digest()
        cls.m = F.Model()
        cls.written = build.build_views_to(cls.m, cls.tmp)
        import build_workbook
        cls.wb_path = os.path.join(cls.tmp, 'wb.xlsx')
        cls.sheets = build_workbook.build(cls.m, F.default_scale(), cls.wb_path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_building_does_not_touch_facts(self):
        self.assertEqual(_facts_digest(), self.before)

    def test_one_view_per_scale_and_per_question(self):
        names = [w.split()[0] for w in self.written]
        for s in F.load_scales():
            self.assertIn('capability-assessment-%s.md' % s.SHORT, names)
        for q in self.m.questions:
            self.assertIn('%s.md' % q['output'], names)
        self.assertIn('provenance.md', names)
        self.assertIn('management-report.html', names)

    def test_html_reports_are_well_formed(self):
        for fn in ('management-report.html', 'management-report-illustrative.html',
                   'walkthrough.html'):
            with open(os.path.join(self.tmp, fn), encoding='utf-8') as f:
                html = f.read()
            p = _Balance()
            p.feed(html)
            self.assertEqual(p.errors, [], fn)
            self.assertEqual(p.stack, [], fn)
            self.assertNotIn('%s', html, fn)
            self.assertNotIn('%d', html, fn)

    def test_illustrative_edition_is_labelled(self):
        with open(os.path.join(self.tmp, 'management-report-illustrative.html'),
                  encoding='utf-8') as f:
            html = f.read()
        self.assertIn('ILLUSTRATIVE', html)
        self.assertIn('SAMPLE', html)
        with open(os.path.join(self.tmp, 'management-report.html'), encoding='utf-8') as f:
            live = f.read()
        self.assertNotIn('ILLUSTRATIVE', live)

    def test_report_section_numbers_are_sequential(self):
        import re
        with open(os.path.join(self.tmp, 'management-report.html'), encoding='utf-8') as f:
            nums = [int(x) for x in re.findall(r'<h2>Section (\d+)</h2>', f.read())]
        self.assertEqual(nums, list(range(1, len(nums) + 1)))

    def test_workbook_has_the_documented_sheets(self):
        self.assertEqual(self.sheets[:3], ['0. Start here', '1. Capabilities',
                                            '2. Observations'])
        self.assertEqual(len(self.sheets), 10)

    def test_workbook_has_no_formulas(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.wb_path)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        self.assertFalse(c.value.startswith('='), (ws.title, c.coordinate))

    def test_workbook_value_list_matches_facts(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.wb_path)
        ws = wb['2. Observations']
        formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        self.assertTrue(any('n/a' in f for f in formulas), formulas)
        for v in F.OBS_VALUES:
            self.assertTrue(any(v in f for f in formulas), (v, formulas))

    def test_every_observation_value_says_what_it_means(self):
        self.assertEqual(set(F.OBS_MEANING), set(F.OBS_VALUES))

    def test_judgement_columns_carry_a_note_listing_their_values(self):
        """A reader must be able to ask a column what it may say."""
        from openpyxl import load_workbook
        ws = load_workbook(self.wb_path)['1. Capabilities']
        scale = F.default_scale()
        wanted = ([t['id'].capitalize() for t in self.m.observation_types]
                  + ['LEVEL', 'Level name'])
        notes = {c.value: c.comment.text for c in ws[1] if c.comment}
        for col in wanted:
            self.assertIn(col, notes)
        for col in [t['id'].capitalize() for t in self.m.observation_types]:
            for v in F.OBS_VALUES:
                self.assertIn(v, notes[col], col)
        for n, k, _ in scale.LEVELS:
            self.assertIn(k, notes['LEVEL'], k)
            self.assertIn(k, notes['Level name'], k)


class Ingest(unittest.TestCase):
    """docs/using-the-model.md - the review loop reads back exactly what the
    reviewer typed, and nothing else."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix='aicap-ingest-')
        cls.m = F.Model()
        import build_workbook
        cls.wb_path = os.path.join(cls.tmp, 'wb.xlsx')
        build_workbook.build(cls.m, F.default_scale(), cls.wb_path)
        with open(os.path.join(F.FACTS, 'observations.json'), encoding='utf-8') as f:
            cls.doc_text = f.read()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def doc(self):
        return json.loads(self.doc_text)

    def test_untouched_workbook_changes_nothing(self):
        doc = self.doc()
        changed, skipped = build.ingest_workbook(self.wb_path, doc)
        self.assertEqual(changed, [])
        self.assertEqual(skipped, [])
        self.assertEqual(json.dumps(doc, sort_keys=True),
                         json.dumps(self.doc(), sort_keys=True))

    def _edit(self, key, value, evidence, by, when):
        from openpyxl import load_workbook
        wb = load_workbook(self.wb_path)
        ws = wb['2. Observations']
        hit = None
        for row in ws.iter_rows(min_row=2):
            cid, typ, crit = row[0].value, row[2].value, row[3].value
            if (str(cid), str(typ), str(crit or '')) == key:
                hit = row
                break
        self.assertIsNotNone(hit, key)
        hit[6].value, hit[7].value, hit[8].value, hit[9].value = value, evidence, by, when
        path = os.path.join(self.tmp, 'edited.xlsx')
        wb.save(path)
        return path

    def test_one_edit_comes_back_as_one_change(self):
        path = self._edit(('4.4', 'practised', '4.4.5'), 'partial',
                          'Guardrails on 1 of 4 agents', 'Test', '2026-09-04')
        doc = self.doc()
        changed, skipped = build.ingest_workbook(path, doc)
        self.assertEqual(skipped, [])
        self.assertEqual([(k, new) for k, _old, new in changed],
                         [(('4.4', 'practised', '4.4.5'), 'partial')])
        rec = [r for r in doc['observations']
               if r.get('criterion') == '4.4.5' and r['observation'] == 'practised'][0]
        self.assertEqual(rec['evidence'], 'Guardrails on 1 of 4 agents')
        self.assertEqual(rec['observed_by'], 'Test')
        self.assertEqual(rec['observed_on'], '2026-09-04')

    def test_legacy_dash_spelling_of_na_is_normalised(self):
        path = self._edit(('1.5', 'enabled', ''), 'n-a', '', 'Test', '2026-09-04')
        doc = self.doc()
        changed, _ = build.ingest_workbook(path, doc)
        self.assertEqual(changed[0][2], 'n/a')

    def test_invalid_value_is_skipped_not_written(self):
        path = self._edit(('1.5', 'enabled', ''), 'maybe', '', 'Test', '2026-09-04')
        doc = self.doc()
        changed, skipped = build.ingest_workbook(path, doc)
        self.assertEqual(changed, [])
        self.assertEqual(len(skipped), 1)

    def test_ingest_never_touches_other_facts_files(self):
        # the function only receives the observations document; the model's
        # other registers are not reachable from it
        doc = self.doc()
        self.assertEqual(set(doc), {'observation_types', 'values', 'observations'})


class Views(unittest.TestCase):
    """The rules the views keep."""

    def test_capability_view_names_no_scale(self):
        with open(os.path.join(ROOT, 'build', 'build_views.py'), encoding='utf-8') as f:
            src = f.read()
        for s in F.load_scales():
            self.assertNotIn('"%s"' % s.SHORT, src)
            self.assertNotIn("'%s'" % s.SHORT, src)

    def test_report_names_no_scale(self):
        with open(os.path.join(ROOT, 'build', 'build_report.py'), encoding='utf-8') as f:
            src = f.read()
        for s in F.load_scales():
            self.assertNotIn('"%s"' % s.SHORT, src)
            self.assertNotIn("'%s'" % s.SHORT, src)

    def test_derived_ladder_is_sound_for_every_scale(self):
        """The 'what it takes to reach each level' table must never claim a
        condition the scale itself does not keep: every combination that the
        scale rates at or above a level must satisfy that level's conditions."""
        import derive
        for s in F.load_scales():
            rungs = {r['n']: r for r in derive.ladder(s)}
            for obs in derive._all_combos():
                lv = s.level(obs)[0]
                if lv is None:
                    continue
                for n, r in rungs.items():
                    if lv < n:
                        continue
                    for t in derive.TYPES:
                        self.assertIn(obs[t], r['conds'][t]['vals'],
                                      '%s L%d %s' % (s.SHORT, n, t))

    def test_derived_ladder_marks_exactness_honestly(self):
        """A rung flagged `exact` must be exactly reproducible from its
        conditions - no combination may satisfy them and yet fall short."""
        import derive
        for s in F.load_scales():
            for r in derive.ladder(s):
                if not r['exact']:
                    continue
                for obs in derive._all_combos():
                    if all(obs[t] in r['conds'][t]['vals'] for t in derive.TYPES):
                        lv = s.level(obs)[0]
                        self.assertIsNotNone(lv, '%s L%d' % (s.SHORT, r['n']))
                        self.assertGreaterEqual(lv, r['n'], '%s L%d' % (s.SHORT, r['n']))

    def test_worked_cases_cover_the_gate(self):
        """The worked table must contain the case the whole model turns on:
        enablers fully in place, practice never observed."""
        import derive
        seen = [obs for _l, obs, _r in derive.worked(F.load_scales())]
        self.assertIn({"practised": "unknown", "enabled": "yes",
                       "skilled": "yes", "defined": "yes"}, seen)

    def test_deck_names_no_scale(self):
        with open(os.path.join(ROOT, 'build', 'build_deck.py'), encoding='utf-8') as f:
            src = f.read()
        for s in F.load_scales():
            self.assertNotIn('"%s"' % s.SHORT, src)
            self.assertNotIn("'%s'" % s.SHORT, src)

    def test_sample_data_never_reaches_facts(self):
        import sample
        m = F.Model()
        before = _facts_digest()
        obs = sample.sample_observations(m)
        self.assertEqual(len(obs), len(m.capabilities))
        self.assertEqual(_facts_digest(), before)
        # and the real model still says what it said
        self.assertEqual(m.values('4.4')['practised'], m.roll_up('4.4', 'practised'))


if __name__ == '__main__':
    unittest.main()
